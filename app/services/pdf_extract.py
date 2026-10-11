"""PDF / DOCX -> plain text. Scanned PDFs are rejected in the MVP (OCR comes later)."""
import io
from dataclasses import dataclass, field


class ExtractionError(Exception):
    pass


class ScannedPDFError(ExtractionError):
    """No text layer found -- most likely an image-only PDF."""


@dataclass
class ExtractedText:
    text: str
    pages: int = 1
    warnings: list[str] = field(default_factory=list)
    links: list[str] = field(default_factory=list)   # real hyperlink targets (not visible in text)
    layout: dict = field(default_factory=dict)       # {'columns_suspected': bool, 'gap_line_ratio': float, 'images': int}


MIN_CHARS = 80


def analyse_layout(pages) -> dict:
    """Heuristic layout check from word positions. A line whose words are split by a wide gap near the middle of
    the page suggests two columns (many ATS read columns in the wrong order). Also counts embedded images."""
    lines = gap_lines = images = 0
    for page in pages:
        images += len(page.images or [])
        width = float(page.width or 1)
        rows: dict[int, list] = {}
        for w in page.extract_words() or []:
            rows.setdefault(round(float(w["top"]) / 3), []).append(w)
        for ws in rows.values():
            if len(ws) < 2:
                continue
            ws.sort(key=lambda w: float(w["x0"]))
            lines += 1
            for a, b in zip(ws, ws[1:]):
                gap = float(b["x0"]) - float(a["x1"])
                mid = (float(a["x1"]) + float(b["x0"])) / 2 / width
                if gap > 0.08 * width and 0.25 <= mid <= 0.75:
                    gap_lines += 1
                    break
    ratio = gap_lines / lines if lines else 0.0
    return {"columns_suspected": gap_lines >= 8 and ratio >= 0.35, "gap_line_ratio": round(ratio, 3), "images": images}


def extract_pdf(data: bytes) -> ExtractedText:
    import pdfplumber

    chunks, warnings, links = [], [], []
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for i, page in enumerate(pdf.pages, 1):
                t = page.extract_text() or ""
                if not t.strip():
                    warnings.append(f"Page {i} has no extractable text")
                chunks.append(t)
                for h in (page.hyperlinks or []):
                    uri = (h.get("uri") or "").strip()
                    if uri.lower().startswith("http") and uri not in links:
                        links.append(uri)
            pages = len(pdf.pages)
            layout = analyse_layout(pdf.pages)
    except Exception as e:  # corrupted / encrypted
        raise ExtractionError(f"Could not read PDF: {e}") from e

    text = "\n".join(chunks).strip()
    if len(text) < MIN_CHARS:
        raise ScannedPDFError(
            "No readable text found. Please upload a text-based PDF (not a scanned image)."
        )
    return ExtractedText(text=text, pages=pages, warnings=warnings, links=links, layout=layout)


def extract_docx(data: bytes) -> ExtractedText:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as e:
        raise ExtractionError(f"Could not read DOCX: {e}") from e
    lines = [p.text for p in document.paragraphs]
    warnings = []
    if document.tables:
        warnings.append("Resume contains tables; text order may be off. Please review.")
        for table in document.tables:
            for row in table.rows:
                lines.append(" | ".join(c.text.strip() for c in row.cells))
    text = "\n".join(lines).strip()
    if len(text) < MIN_CHARS:
        raise ExtractionError("DOCX has too little text to parse.")
    links = [r.target_ref for r in document.part.rels.values()
             if r.reltype.endswith("/hyperlink") and r.is_external
             and r.target_ref.lower().startswith("http")]
    return ExtractedText(text=text, warnings=warnings, links=list(dict.fromkeys(links)))


def extract_text(filename: str, data: bytes) -> ExtractedText:
    name = filename.lower()
    if name.endswith(".pdf"):
        return extract_pdf(data)
    if name.endswith(".docx"):
        return extract_docx(data)
    raise ExtractionError("Unsupported file type. Upload a .pdf or .docx file.")
